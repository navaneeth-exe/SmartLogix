import os
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX_PATHS = [
    r"c:\Users\NAVANEETH\Documents\Academic Projects\Smart Inventory & Delivery Optimization System\Project Report\SmartLogix_Project_Report_Formatted.docx",
    r"c:\Users\NAVANEETH\Documents\Academic Projects\Smart Inventory & Delivery Optimization System\Project Report\SmartLogix_Project_Report_Final.docx",
    r"c:\Users\NAVANEETH\Documents\Academic Projects\Smart_Inventory_and_Delivery_Optimization_System\Project Report\SmartLogix_Project_Report_Formatted.docx",
    r"c:\Users\NAVANEETH\Documents\Academic Projects\Smart_Inventory_and_Delivery_Optimization_System\Project Report\SmartLogix_Project_Report_Final.docx",
]

MD_PATH = r"c:\Users\NAVANEETH\Documents\Academic Projects\Smart_Inventory_and_Delivery_Optimization_System\Project Report\SMARTLOGIX_PROJECT_REPORT.md"

# Structured segments for Acknowledgement with bold names
# Format: (text, is_bold)
PARAGRAPHS = [
    [
        ("First and foremost, we would like to thank the ", False),
        ("GOD ALMIGHTY", True),
        (" for his infinite grace and help, without which this project would not have reached its completion.", False),
    ],
    [
        ("We wish to express our sincere gratitude to ", False),
        ("Dr. P R Suresh, Ph.D.", True),
        (", Principal, Ahalia School of Engineering and Technology for his timely support throughout the course of this project.", False),
    ],
    [
        ("We extend our sincere thanks to ", False),
        ("Dr. Krishna Kumar Kishor, Ph.D.", True),
        (", Vice Principal, for extending all facilities.", False),
    ],
    [
        ("We express our sincere and heartfelt gratitude to ", False),
        ("Dr. S. Gunasekaran, Ph.D.", True),
        (", Professor and Head of the Department, Department of Computer Science and Engineering for his help and excellent encouragements through the project.", False),
    ],
    [
        ("We express our deepest, heartiest and sincere thanks to our Project Guide ", False),
        ("Ms. Happy A", True),
        (", Assistant Professor, for her exhilarating supervision, timely suggestions and encouragement during all phases of this work.", False),
    ],
    [
        ("We take this opportunity to thank ", False),
        ("Dr. Reshmi B", True),
        (", our mentor who had been always there with us.", False),
    ],
    [
        ("Also, we would like to thank all the Teaching faculty members and Supporting Staff of our department.", False),
    ],
    [
        ("We would like to convey our gratitude to our ", False),
        ("Parents", True),
        (" whose prayers and blessings were always with us.", False),
    ],
    [
        ("Last but not least we would like to thank our ", False),
        ("Friends", True),
        (" and others who directly or indirectly helped us in successful completion of this work.", False),
    ],
]

def update_docx(path):
    if not os.path.exists(path):
        return
    print(f"Updating DOCX: {path}")
    doc = docx.Document(path)

    # 1. Update Acknowledgement paragraph
    found_ack = False
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "ACKNOWLEDGEMENT":
            for j in range(i + 1, min(i + 5, len(doc.paragraphs))):
                cand = doc.paragraphs[j]
                if cand.text.strip():
                    cand.clear()
                    cand.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    cand.paragraph_format.line_spacing = 1.5
                    cand.paragraph_format.space_after = Pt(12)

                    first_para = True
                    for para_segs in PARAGRAPHS:
                        if not first_para:
                            r_nl = cand.add_run("\n\n")
                            r_nl.font.name = "Times New Roman"
                            r_nl.font.size = Pt(12)
                        first_para = False

                        for seg_text, is_bold in para_segs:
                            run = cand.add_run(seg_text)
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(12)
                            run.font.bold = is_bold
                            run.font.color.rgb = RGBColor(0, 0, 0)
                    found_ack = True
                    print(f"  Successfully updated ACKNOWLEDGEMENT in paragraph {j}!")
                    break
            if found_ack:
                break

    # 2. Update Guide name from Ms Deepthy P S to Ms. Happy A across Certificate, Declaration, Table 1
    for p in doc.paragraphs:
        if "Ms Deepthy P S" in p.text:
            for r in p.runs:
                if "Ms Deepthy P S" in r.text:
                    r.text = r.text.replace("Ms Deepthy P S", "Ms. Happy A")
            if "Ms Deepthy P S" in p.text:
                full_text = p.text.replace("Ms Deepthy P S", "Ms. Happy A")
                p.text = full_text

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if "Ms Deepthy P S" in cell.text:
                    cell.text = cell.text.replace("Ms Deepthy P S", "Ms. Happy A")
                    for cp in cell.paragraphs:
                        for cr in cp.runs:
                            cr.font.name = "Times New Roman"
                            cr.font.size = Pt(12)

    doc.save(path)
    print(f"  Saved {path}")

def update_markdown(path):
    if not os.path.exists(path):
        return
    print(f"Updating Markdown: {path}")
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()

    text = text.replace("Ms Deepthy P S", "Ms. Happy A")

    ack_header = "## ACKNOWLEDGEMENT\n"
    ack_end = "\n---\n\n## INSTITUTION & DEPARTMENT VISION AND MISSION"
    
    start_pos = text.find(ack_header)
    end_pos = text.find(ack_end, start_pos)

    new_ack_md = """## ACKNOWLEDGEMENT
First and foremost, we would like to thank the **GOD ALMIGHTY** for his infinite grace and help, without which this project would not have reached its completion.  
We wish to express our sincere gratitude to **Dr. P R Suresh, Ph.D.**, Principal, Ahalia School of Engineering and Technology for his timely support throughout the course of this project.  
We extend our sincere thanks to **Dr. Krishna Kumar Kishor, Ph.D.**, Vice Principal, for extending all facilities.  
We express our sincere and heartfelt gratitude to **Dr. S. Gunasekaran, Ph.D.**, Professor and Head of the Department, Department of Computer Science and Engineering for his help and excellent encouragements through the project.  
We express our deepest, heartiest and sincere thanks to our Project Guide **Ms. Happy A**, Assistant Professor, for her exhilarating supervision, timely suggestions and encouragement during all phases of this work.  
We take this opportunity to thank **Dr. Reshmi B**, our mentor who had been always there with us.  
Also, we would like to thank all the Teaching faculty members and Supporting Staff of our department.  
We would like to convey our gratitude to our **Parents** whose prayers and blessings were always with us.  
Last but not least we would like to thank our **Friends** and others who directly or indirectly helped us in successful completion of this work."""

    if start_pos != -1 and end_pos != -1:
        text = text[:start_pos] + new_ack_md + text[end_pos:]
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
        print("  Successfully updated SMARTLOGIX_PROJECT_REPORT.md!")
    else:
        print("  Could not locate acknowledgement boundaries in markdown.")

if __name__ == "__main__":
    for p in DOCX_PATHS:
        update_docx(p)
    update_markdown(MD_PATH)
    print("\nAll documents updated successfully!")
