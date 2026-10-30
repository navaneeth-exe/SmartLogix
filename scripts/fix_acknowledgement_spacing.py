import os
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_PATHS = [
    r"c:\Users\NAVANEETH\Documents\Academic Projects\Smart Inventory & Delivery Optimization System\Project Report\SmartLogix_Project_Report_Formatted.docx",
    r"c:\Users\NAVANEETH\Documents\Academic Projects\Smart Inventory & Delivery Optimization System\Project Report\SmartLogix_Project_Report_Final.docx",
]

# 9 paragraphs for Acknowledgement
ACK_PARAS = [
    [
        ("First and foremost, we would like to thank the ", False),
        ("GOD ALMIGHTY", True),
        (" for his infinite grace and help, without which this project would not have reached its completion.", False),
    ],
    [
        ("We wish to express our sincere gratitude to ", False),
        ("Dr. P R Suresh, Ph.D.", True),
        (", Principal, Ahalia School of Engineering and Technology for his timely support throughout the course of this project.", False),
    ],
    [
        ("We extend our sincere thanks to ", False),
        ("Dr. Krishna Kumar Kishor, Ph.D.", True),
        (", Vice Principal, for extending all facilities.", False),
    ],
    [
        ("We express our sincere and heartfelt gratitude to ", False),
        ("Dr. S. Gunasekaran, Ph.D.", True),
        (", Professor and Head of the Department, Department of Computer Science and Engineering for his help and excellent encouragements through the project.", False),
    ],
    [
        ("We express our deepest, heartiest and sincere thanks to our Project Guide ", False),
        ("Ms. Happy A", True),
        (", Assistant Professor, for her exhilarating supervision, timely suggestions and encouragement during all phases of this work.", False),
    ],
    [
        ("We take this opportunity to thank ", False),
        ("Dr. Reshmi B", True),
        (", our mentor who had been always there with us.", False),
    ],
    [
        ("Also, we would like to thank all the Teaching faculty members and Supporting Staff of our department.", False),
    ],
    [
        ("We would like to convey our gratitude to our ", False),
        ("Parents", True),
        (" whose prayers and blessings were always with us.", False),
    ],
    [
        ("Last but not least we would like to thank our ", False),
        ("Friends", True),
        (" and others who directly or indirectly helped us in successful completion of this work.", False),
    ],
]

def make_clean_para(doc, segs):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    for seg_text, is_bold in segs:
        run = p.add_run(seg_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = is_bold
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def fix_docx(path):
    print(f"\nProcessing {path}...")
    doc = docx.Document(path)

    # 1. Replace the single giant P30 paragraph under ACKNOWLEDGEMENT with 9 distinct paragraphs
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "ACKNOWLEDGEMENT":
            target_p = doc.paragraphs[i + 1]
            target_elem = target_p._p
            for segs in ACK_PARAS:
                new_p = make_clean_para(doc, segs)
                target_elem.addprevious(new_p._p)
            target_elem.getparent().remove(target_elem)
            print("  Replaced single Acknowledgement paragraph with 9 clean distinct paragraphs!")
            break

    # 2. Split itemized MISSION paragraphs into clean individual paragraphs if they contain newlines
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t in ["MISSION OF THE INSTITUTION", "MISSION OF THE DEPARTMENT (CSE)", "MISSION OF THE DEPARTMENT"]:
            target_p = doc.paragraphs[i + 1]
            if "\n" in target_p.text:
                target_elem = target_p._p
                lines = [l.strip() for l in target_p.text.split("\n") if l.strip()]
                for line in lines:
                    new_p = doc.add_paragraph()
                    new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    new_p.paragraph_format.line_spacing = 1.5
                    new_p.paragraph_format.space_before = Pt(0)
                    new_p.paragraph_format.space_after = Pt(4)
                    r = new_p.add_run(line)
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    target_elem.addprevious(new_p._p)
                target_elem.getparent().remove(target_elem)
                print(f"  Split {t} into {len(lines)} clean paragraphs!")

    # 3. Split ABSTRACT paragraphs if newlines exist
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "ABSTRACT":
            target_p = doc.paragraphs[i + 1]
            if "\n" in target_p.text:
                target_elem = target_p._p
                blocks = [b.strip() for b in target_p.text.split("\n") if b.strip()]
                for b in blocks:
                    new_p = doc.add_paragraph()
                    new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    new_p.paragraph_format.line_spacing = 1.5
                    new_p.paragraph_format.space_before = Pt(0)
                    new_p.paragraph_format.space_after = Pt(6)
                    r = new_p.add_run(b)
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    target_elem.addprevious(new_p._p)
                target_elem.getparent().remove(target_elem)
                print("  Split ABSTRACT into clean paragraphs!")
            break

    doc.save(path)
    print(f"  Saved {path}")

if __name__ == "__main__":
    for p in DOCX_PATHS:
        fix_docx(p)
    print("\nSuccessfully finished fixing formatting!")
